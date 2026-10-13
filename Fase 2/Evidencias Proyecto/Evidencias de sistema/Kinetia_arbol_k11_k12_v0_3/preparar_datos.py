"""Extrae el Excel fijado sin modificarlo ni completar etiquetas."""
from datetime import datetime
from base import (BASE, CAMPOS, CONTEXTO, SOURCE_SHA, exigir, guardar, sha,
                  leer, verificar_fuentes, validar_datos)

def extraer_original():
    import openpyxl
    ruta = BASE / "fuentes/revision_completada.xlsx"
    exigir(sha(ruta) == SOURCE_SHA, "No es la fuente fijada.")
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=False)
    rev, raw = wb["Revisar"], wb["Datos y criterios"]
    casos = []
    for i in range(8):
        rr, dr = 11 + i, 7 + i
        cid = rev.cell(rr, 1).value
        exigir(cid == raw.cell(dr, 1).value == f"DB{i+1:02}", "Filas desalineadas.")
        aceptables = rev.cell(rr, 4).value
        exigir(aceptables in ["AMBAS", "K11", "K12"], "Aceptabilidad pendiente.")
        casos.append({
            "caso_id": cid, "grupo_id": raw.cell(dr, 2).value,
            "caso_origen": raw.cell(dr, 3).value, "escenario_origen": raw.cell(dr, 10).value,
            "origen": raw.cell(dr, 4).value,
            "historia": rev.cell(rr, 2).value,
            "entradas": dict(zip(CAMPOS, [raw.cell(dr, c).value for c in [5,6,7,8]])),
            "meses_exactos_conocidos": raw.cell(dr, 9).value,
            "revision": {
                "preferencia": rev.cell(rr, 3).value,
                "aceptables": ["K11", "K12"] if aceptables == "AMBAS" else [aceptables],
                "firmeza": rev.cell(rr, 5).value, "motivo": rev.cell(rr, 6).value,
                "procedencia_texto_original": rev.cell(rr, 7).value,
                "rango_excel": f"Revisar!A{rr}:G{rr}",
            }
        })
    fecha = rev["F8"].value
    exigir(isinstance(fecha, datetime), "La fecha del Excel no es una fecha.")
    datos = {
        "version": "0.1", "uso": "desarrollo", "contexto": CONTEXTO,
        "fuente": {"archivo": ruta.name, "sha256": SOURCE_SHA,
                  "fecha_recepcion_chile": "2026-10-10",
                  "fecha_revision_declarada": fecha.date().isoformat()},
        "revisor": {"identificacion_declarada": rev["B8"].value, "rol_declarado": rev["D8"].value,
                    "identidad_estable_verificada": False, "credenciales_verificadas": False},
        "notas": [
            "Los textos de pendiente G13/G16 y B29 no anulan las respuestas completas C:F.",
            "B40 conserva una nota anterior. Identificacion generica; no se atribuye nombre.",
            "Ocho derivados de G03, no ocho observaciones independientes.",
            "DB01 representa C04: no agregar C04 nuevamente.",
            "La dosis pasada no se conoce y no se exige que coincida con la actual.",
            "Meses exactos conocidos es procedencia, no un quinto predictor.",
        ],
        "casos": casos,
    }
    wb.close()
    return datos

def extraer():
    verificar_fuentes()
    datos = extraer_original()
    chat = leer(BASE / "fuentes/respuestas_DB09_DB11.json")
    nuevos = chat["casos"]
    exigir([c["caso_id"] for c in nuevos] == ["DB09", "DB10", "DB11"],
           "Las respuestas nuevas deben conservar DB09-DB11 en orden.")
    datos["version"] = "0.2"
    datos["fuente_chat"] = {k: v for k, v in chat.items() if k != "casos"}
    datos["casos"].extend(nuevos)
    datos["notas"].extend([
        "Las once filas combinadas siguen en G03; conservar tambien los parientes de K06/K08 en ese grupo.",
        "DB09-DB11 amplian desarrollo tras examinar discrepancias de v0.1; no son reserva.",
        "DB11 EMPATE se conserva como control descriptivo; no entra en fit ni en coincidencia de preferencia unica.",
        "DB01-DB10 aportan una fila por caso, sin duplicacion ni pesos de empate.",
    ])
    return validar_datos(datos)

if __name__ == "__main__":
    salida = BASE / "datos/desarrollo_k11_k12.json"
    guardar(salida, extraer())
    print("Extraccion completada: once casos G03, diez con preferencia y un empate; fuentes intactas.")

